#!/usr/bin/env python3
"""Download official US data (FRED, Census) and write data/data.json for the site.

Usage:  pip install openpyxl && python3 scripts/build_data.py
All figures are annual (calendar-year averages of monthly/weekly/quarterly data).
"""
import csv
import io
import json
import os
import re
import time
import urllib.request
from collections import defaultdict
from datetime import date

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "data", "data.json")

FRED = {
    # key: (series id, description)
    "homePrice": ("MSPUS", "Median sales price of houses sold, US (Census/HUD)"),
    "mortgageRate": ("MORTGAGE30US", "30-year fixed mortgage rate, % (Freddie Mac)"),
    "familyIncome": ("MEFAINUSA646N", "Median family income, current $ (Census)"),
    "cpi": ("CPIAUCSL", "CPI-U all items (BLS)"),
    "rentIdx": ("CUUR0000SEHA", "CPI: rent of primary residence (BLS)"),
    "tuitionIdx": ("CUUR0000SEEB", "CPI: tuition, other school fees & childcare (BLS)"),
    "medicalIdx": ("CPIMEDSL", "CPI: medical care (BLS)"),
    "foodIdx": ("CPIUFDSL", "CPI: food (BLS)"),
    "carIdx": ("CUUR0000SETA01", "CPI: new vehicles (BLS)"),
    "gasIdx": ("CUUR0000SETB01", "CPI: gasoline (BLS)"),
    "apparelIdx": ("CPIAPPSL", "CPI: apparel (BLS)"),
    "savingRate": ("PSAVERT", "Personal saving rate, % of disposable income (BEA)"),
    "ownershipAll": ("RHORUSQ156N", "Homeownership rate, all ages, % (Census HVS)"),
    "minWage": ("FEDMINNFRWG", "Federal minimum wage, $/hr (DOL)"),
    "productivity": ("OPHNFB", "Output per hour, nonfarm business (BLS)"),
    "realComp": ("COMPRNFB", "Real hourly compensation, nonfarm business (BLS)"),
    "hourlyWage": ("AHETPI", "Avg hourly earnings, production & nonsupervisory workers, $ (BLS)"),
    "studentDebt": ("SLOAS", "Student loans owned & securitized, $ millions (Fed)"),
}

CENSUS_INCOME = "https://www2.census.gov/programs-surveys/cps/tables/time-series/historical-income-households/"
HVS_AGE = "https://www.census.gov/housing/hvs/data/histtab19.xlsx"
RENT_DECENNIAL = "https://www2.census.gov/programs-surveys/decennial/tables/time-series/coh-grossrents/grossrents-unadj.txt"


def get(url, tries=4):
    req = urllib.request.Request(url, headers={"User-Agent": "curl/8.5.0", "Accept": "*/*"})
    for i in range(tries):
        try:
            with urllib.request.urlopen(req, timeout=60) as r:
                return r.read()
        except Exception:
            if i == tries - 1:
                raise
            time.sleep(2 ** i)


def fred_annual(series):
    text = get(f"https://fred.stlouisfed.org/graph/fredgraph.csv?id={series}").decode()
    buckets = defaultdict(list)
    for row in csv.DictReader(io.StringIO(text)):
        v = row[series]
        if v in ("", "."):
            continue
        buckets[int(row["observation_date"][:4])].append(float(v))
    return {y: sum(v) / len(v) for y, v in buckets.items()}, {y: len(v) for y, v in buckets.items()}


def year_of(cell):
    m = re.match(r"^(\d{4})", str(cell or ""))
    return int(m.group(1)) if m else None


def census_income_block(wb_bytes, header_label, value_col=2):
    """Parse a Census H-table block: rows under `header_label`, first occurrence of each year wins
    (Census lists the newest methodology first when a year is repeated)."""
    ws = openpyxl.load_workbook(io.BytesIO(wb_bytes), data_only=True).worksheets[0]
    out, inside = {}, False
    for row in ws.iter_rows(values_only=True):
        a = row[0]
        if isinstance(a, str) and a.strip() == header_label:
            inside = True
            continue
        if not inside:
            continue
        y = year_of(a)
        if y is None:
            if out and isinstance(a, str) and a.strip() and not a.startswith(("Age", "Race", "Current")):
                break  # next block begins
            continue
        if y not in out and isinstance(row[value_col], (int, float)):
            out[y] = float(row[value_col])
    return out


def hvs_under35():
    ws = openpyxl.load_workbook(io.BytesIO(get(HVS_AGE)), data_only=True).worksheets[0]
    buckets, cur = defaultdict(list), None
    for row in ws.iter_rows(values_only=True):
        a = row[0]
        y = year_of(a)
        if y and all(c is None for c in row[1:]):
            cur = y
            continue
        if cur and isinstance(a, str) and re.match(r"^\s*(1st|2nd|3rd|4th)", a) and isinstance(row[2], (int, float)):
            buckets[cur].append(float(row[2]))
    return {y: sum(v) / len(v) for y, v in buckets.items()}


def decennial_rent():
    text = get(RENT_DECENNIAL).decode()
    lines = text.splitlines()
    years = [int(x) for x in re.findall(r"\d{4}", next(l for l in lines if "2000" in l and "1990" in l))]
    us = next(l for l in lines if l.startswith("United States"))
    vals = [int(x) for x in re.findall(r"\$(\d+)", us)]
    return dict(zip(years, vals))


def main():
    series, notes = {}, {}
    for key, (sid, desc) in FRED.items():
        vals, counts = fred_annual(sid)
        series[key] = vals
        notes[key] = {"source": desc, "id": sid, "url": f"https://fred.stlouisfed.org/series/{sid}",
                      "partial": {y: c for y, c in counts.items() if y == max(counts)}}
        print(f"{key:14s} {sid:16s} {min(vals)}-{max(vals)}")

    h05 = get(CENSUS_INCOME + "h05.xlsx")
    series["hhIncome"] = census_income_block(h05, "All Races")
    notes["hhIncome"] = {"source": "Median household income, current $ (Census Table H-5)",
                         "url": CENSUS_INCOME + "h05.xlsx"}
    h10 = get(CENSUS_INCOME + "h10ar.xlsx")
    series["youngIncome"] = census_income_block(h10, "25 to 34 Years")
    notes["youngIncome"] = {"source": "Median income of households headed by 25-34 year olds, current $ (Census Table H-10)",
                            "url": CENSUS_INCOME + "h10ar.xlsx"}
    series["ownershipUnder35"] = hvs_under35()
    notes["ownershipUnder35"] = {"source": "Homeownership rate, householders under 35, % (Census HVS Table 19)", "url": HVS_AGE}
    print("hhIncome", min(series["hhIncome"]), max(series["hhIncome"]))
    print("youngIncome", min(series["youngIncome"]), max(series["youngIncome"]))
    print("under35", min(series["ownershipUnder35"]), max(series["ownershipUnder35"]))

    # Median gross rent in dollars: Census decennial values (1960-2000), with years in between
    # interpolated along the CPI rent index, and years after 2000 extended with the same index.
    anchors = decennial_rent()
    idx = series["rentIdx"]
    rent = {}
    ys = sorted(anchors)
    for a, b in zip(ys, ys[1:]):
        if a < 1960:
            continue
        # scale factor drifts linearly between the two anchors so both are hit exactly
        fa, fb = anchors[a] / idx[a], anchors[b] / idx[b]
        for y in range(a, b + 1):
            t = (y - a) / (b - a)
            rent[y] = idx[y] * (fa + (fb - fa) * t)
    f2000 = anchors[2000] / idx[2000]
    for y in idx:
        if y > 2000:
            rent[y] = idx[y] * f2000
    series["rent"] = rent
    notes["rent"] = {"source": "Median gross rent, $/month: Census decennial census 1960-2000 "
                               "(interpolated/extended with the CPI rent index; an estimate between census years)",
                     "url": "https://www.census.gov/data/tables/time-series/dec/coh-grossrents.html",
                     "anchors": anchors}

    all_years = sorted({y for s in series.values() for y in s if y >= 1960})
    out = {
        "generated": date.today().isoformat(),
        "years": all_years,
        "series": {k: [round(v[y], 4) if y in v else None for y in all_years] for k, v in series.items()},
        "notes": notes,
    }
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    with open(OUT, "w") as f:
        json.dump(out, f, separators=(",", ":"))
    print("wrote", OUT, os.path.getsize(OUT), "bytes")


if __name__ == "__main__":
    main()
